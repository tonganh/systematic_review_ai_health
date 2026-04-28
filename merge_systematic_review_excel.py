from __future__ import annotations

import difflib
import re
import shutil
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


def normalize_text(text: str) -> str:
    value = text.strip().replace("–", "-").replace("—", "-")
    value = value.replace('"', "").replace("“", "").replace("”", "").replace("'", "")
    value = unicodedata.normalize("NFKD", value)
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    value = value.lower()
    value = re.sub(r"\s*-\s*[^-]*$", "", value)
    value = re.sub(r"[^a-z0-9]+", "", value)
    return value


def main() -> None:
    root = Path(__file__).resolve().parent
    target_path = root / "outputs" / "Systematic review.xlsx"
    source_path = root / "outputs" / "review_summary.xlsx"
    backup_path = root / "outputs" / "Systematic review.backup.xlsx"

    shutil.copy2(target_path, backup_path)

    target_wb = load_workbook(target_path)
    source_wb = load_workbook(source_path)
    target_ws = target_wb[target_wb.sheetnames[0]]
    source_ws = source_wb["papers_answers"]

    source_headers = [cell.value for cell in next(source_ws.iter_rows(min_row=1, max_row=1))]
    source_rows = []
    for row in source_ws.iter_rows(min_row=2, max_row=source_ws.max_row, values_only=True):
        record = {source_headers[i]: row[i] for i in range(len(source_headers))}
        title_key = normalize_text(str(record.get("title") or ""))
        if title_key:
            source_rows.append((title_key, record))

    best_by_title: dict[str, dict] = {}
    for key, record in source_rows:
        current = best_by_title.get(key)
        if current is None:
            best_by_title[key] = record
            continue
        source_file = str(record.get("source_file") or "")
        current_file = str(current.get("source_file") or "")
        score = (0 if "(1)" in source_file else 1, len(source_file))
        current_score = (0 if "(1)" in current_file else 1, len(current_file))
        if score > current_score:
            best_by_title[key] = record

    merge_fields = [
        "source_file",
        "title",
        "year",
        "type_of_paper",
        "population",
        "ai_type_application",
        "mental_health_outcome",
        "positive_effects",
        "negative_effects",
        "evidence_type",
        "main_finding",
        "limitations_bias",
        "relevance_to_our_review",
        "use_decision",
        "notes_for_research_question",
        "is_review_paper",
        "primary_studies_count",
    ]
    target_headers = [
        "matched_source_file",
        "matched_title",
        "matched_year",
        "type_of_paper",
        "population",
        "ai_type_application",
        "mental_health_outcome",
        "positive_effects",
        "negative_effects",
        "evidence_type",
        "main_finding",
        "limitations_bias",
        "relevance_to_our_review",
        "use_decision",
        "notes_for_research_question",
        "is_review_paper",
        "primary_studies_count",
        "match_status",
    ]

    existing_headers = [cell.value for cell in next(target_ws.iter_rows(min_row=1, max_row=1))]
    start_col = len(existing_headers) + 1
    for offset, header in enumerate(target_headers):
        target_ws.cell(row=1, column=start_col + offset, value=header)

    available_keys = list(best_by_title.keys())
    matched = 0
    for row_idx in range(2, target_ws.max_row + 1):
        paper_cell = str(target_ws.cell(row=row_idx, column=2).value or "")
        base_title = paper_cell.split("—")[0].strip().strip('"').strip("“").strip("”")
        title_key = normalize_text(base_title)
        record = best_by_title.get(title_key)
        status = "exact"
        if record is None and title_key:
            close = difflib.get_close_matches(title_key, available_keys, n=1, cutoff=0.9)
            if close:
                record = best_by_title[close[0]]
                status = "fuzzy"
        if record is None:
            values = [""] * len(merge_fields) + ["not_found"]
        else:
            matched += 1
            values = [record.get(field, "") for field in merge_fields] + [status]
        for offset, value in enumerate(values):
            target_ws.cell(row=row_idx, column=start_col + offset, value=value)

    target_wb.save(target_path)

    total = target_ws.max_row - 1
    print(f"backup={backup_path}")
    print(f"rows_total={total}")
    print(f"rows_matched={matched}")
    print(f"rows_unmatched={total - matched}")


if __name__ == "__main__":
    main()

